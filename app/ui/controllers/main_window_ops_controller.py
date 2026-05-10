from __future__ import annotations

import json
import os
from datetime import date
from pathlib import Path

from PySide6.QtWidgets import QDialog, QFileDialog, QInputDialog, QMessageBox

from app.audit import local_actor_identifier
from app.ui.license_dialog import LicenseActivationDialog, LicenseManageDialog
from app.utils.backup import create_backup, export_backup, inspect_backup_counts, restore_backup


class MainWindowOpsController:
    """Cross-cutting operational logic delegated from MainWindow."""

    def __init__(self, window) -> None:
        self.window = window

    def log_audit(
        self,
        action_type: str,
        entity_type: str = "",
        entity_id: str = "",
        details: str = "",
    ) -> None:
        w = self.window
        try:
            w.bookkeeping_service.log_audit(
                actor_role=local_actor_identifier(w.current_role),
                action_type=action_type,
                entity_type=entity_type,
                entity_id=entity_id,
                details=details,
            )
        except Exception:
            pass

    def ask_admin_pin(self) -> str | None:
        pin, ok = QInputDialog.getText(self.window, "Admin PIN", "Enter admin PIN:")
        if not ok:
            return None
        return pin.strip()

    def require_admin_access(self, action_name: str) -> str | None:
        pin = self.ask_admin_pin()
        if pin is None:
            return None
        if not self.window.bookkeeping_service.verify_admin_pin(pin):
            QMessageBox.warning(self.window, action_name, "Invalid admin PIN.")
            return None
        return pin

    def ensure_license_activation(self) -> bool:
        w = self.window
        if not w.license_service.is_enforcement_configured():
            return True

        status = w.license_service.evaluate_current_license()
        if status.valid:
            w.license_service.mark_successful_launch()
            return True

        dialog = LicenseActivationDialog(license_service=w.license_service, parent=w)
        if dialog.exec() != QDialog.Accepted:
            return False

        post_status = w.license_service.evaluate_current_license()
        if not post_status.valid:
            QMessageBox.critical(
                w,
                "Activation Failed",
                f"Activation did not complete successfully:\n{post_status.reason}",
            )
            return False

        w.license_service.mark_successful_launch()
        self.log_audit("license_activate", "license", post_status.payload.license_id if post_status.payload else "", post_status.reason)
        return True

    def manage_license(self) -> None:
        w = self.window
        dialog = LicenseManageDialog(license_service=w.license_service, parent=w)
        if dialog.exec() == QDialog.Accepted:
            status = w.license_service.evaluate_current_license()
            if status.valid:
                w.license_service.mark_successful_launch()
                self.log_audit(
                    "license_update",
                    "license",
                    status.payload.license_id if status.payload else "",
                    status.reason,
                )
                w._update_shell_status()

    def on_role_changed(self, role_name: str) -> None:
        w = self.window
        role = role_name.strip().lower()
        if role == w.current_role:
            return

        if role == "admin":
            pin = self.ask_admin_pin()
            if pin is None or not w.bookkeeping_service.verify_admin_pin(pin):
                QMessageBox.warning(w, "Role Switch", "Admin PIN verification failed.")
                w.role_combo.blockSignals(True)
                w.role_combo.setCurrentText(w.current_role)
                w.role_combo.blockSignals(False)
                return

        w.current_role = role
        w.bookkeeping_service.set_setting("current_role", role)
        self.apply_role_permissions()
        w._update_shell_status()
        self.log_audit("role_switch", "session", "current", f"Role changed to {role}")

    def apply_role_permissions(self) -> None:
        # All tabs remain enabled; sensitive actions are protected by PIN prompts.
        return

    def configure_auto_backup_timer(self) -> None:
        w = self.window
        enabled = (w.bookkeeping_service.get_setting("auto_backup_enabled", "0") or "0") == "1"
        interval_str = w.bookkeeping_service.get_setting("backup_interval_minutes", "60") or "60"
        try:
            interval_minutes = max(5, int(float(interval_str)))
        except ValueError:
            interval_minutes = 60

        if hasattr(w, "auto_backup_enabled_checkbox"):
            w.auto_backup_enabled_checkbox.setChecked(enabled)
        if hasattr(w, "auto_backup_interval_spin"):
            w.auto_backup_interval_spin.setValue(interval_minutes)

        w.auto_backup_timer.stop()
        if enabled:
            w.auto_backup_timer.setInterval(interval_minutes * 60 * 1000)
            w.auto_backup_timer.start()

    def save_backup_preferences(self) -> None:
        w = self.window
        enabled = w.auto_backup_enabled_checkbox.isChecked()
        interval = int(w.auto_backup_interval_spin.value())
        w.bookkeeping_service.set_setting("auto_backup_enabled", "1" if enabled else "0")
        w.bookkeeping_service.set_setting("backup_interval_minutes", str(interval))
        self.configure_auto_backup_timer()
        self.log_audit(
            "backup_schedule_update",
            "settings",
            "auto_backup",
            f"enabled={enabled}, interval_minutes={interval}",
        )
        QMessageBox.information(w, "Backup Schedule", "Backup schedule updated.")

    def run_scheduled_backup(self) -> None:
        w = self.window
        try:
            backup_path = create_backup(db_path=w.db_path)
            self.log_audit("scheduled_backup", "backup", str(backup_path), "Automatic backup completed")
        except Exception as exc:
            self.log_audit("scheduled_backup_failed", "backup", "", str(exc))

    @staticmethod
    def format_restore_preview(current_counts: dict, backup_counts: dict, backup_file: str) -> str:
        lines = [
            "Restoring will overwrite current live database.",
            "",
            f"Backup file: {backup_file}",
            "",
            "Current DB -> Backup DB",
            f"Items: {current_counts.get('items', 0)} -> {backup_counts.get('items', 0)}",
            f"Sales: {current_counts.get('sales', 0)} -> {backup_counts.get('sales', 0)}",
            f"Purchases: {current_counts.get('purchases', 0)} -> {backup_counts.get('purchases', 0)}",
            f"Expenses: {current_counts.get('expenses', 0)} -> {backup_counts.get('expenses', 0)}",
            f"Stock Movements: {current_counts.get('stock_movements', 0)} -> {backup_counts.get('stock_movements', 0)}",
            "",
            "Continue with restore?",
        ]
        return "\n".join(lines)

    def export_reports_xlsx(self) -> None:
        w = self.window
        try:
            from openpyxl import Workbook
        except Exception:
            QMessageBox.warning(w, "XLSX Export", "openpyxl is not installed. Run: pip install openpyxl")
            return

        start_date, end_date = w._iso_range_from_edits(w.report_from_date, w.report_to_date)
        path, _ = QFileDialog.getSaveFileName(
            w,
            "Save XLSX",
            f"reports_{w._range_suffix(start_date, end_date)}.xlsx",
            "Excel Workbook (*.xlsx)",
        )
        if not path:
            return

        wb = Workbook()
        ws_summary = wb.active
        ws_summary.title = "Summary"
        ws_summary.append(["Metric", "Value"])
        ws_summary.append(["From", start_date])
        ws_summary.append(["To", end_date])
        ws_summary.append(["Sales", w.sales_value.text().replace("INR ", "")])
        ws_summary.append(["COGS", w.cogs_value.text().replace("INR ", "")])
        ws_summary.append(["Gross Profit", w.gross_profit_value.text().replace("INR ", "")])
        ws_summary.append(["Purchases", w.purchases_value.text().replace("INR ", "")])
        ws_summary.append(["Expenses", w.expenses_value.text().replace("INR ", "")])
        ws_summary.append(["Period Fixed Cost", w.fixed_daily_value.text().replace("INR ", "")])
        ws_summary.append(["Net Profit", w.net_value.text().replace("INR ", "")])
        if hasattr(w, "waste_summary_label"):
            ws_summary.append(["Waste", w.waste_summary_label.text().replace("Waste in range: ", "")])

        table_sheets = [
            ("SalesTrend", w.sales_trend_table),
            ("TopItems", w.top_items_table),
            ("PaymentBreakdown", w.payment_breakdown_table),
            ("RecentSales", w.recent_sales_table),
            ("LowStock", w.low_stock_table),
            ("StockLedger", w.ledger_table),
        ]
        for sheet_name, table in table_sheets:
            ws = wb.create_sheet(title=sheet_name)
            for row in w._table_rows_for_csv(table):
                ws.append(row)

        try:
            wb.save(path)
        except Exception as exc:
            QMessageBox.critical(w, "XLSX Export", f"Failed to save XLSX: {exc}")
            return

        self.log_audit("xlsx_export", "report", "range", f"{start_date}..{end_date}")
        if w.open_after_export_checkbox.isChecked() and hasattr(os, "startfile"):
            try:
                os.startfile(path)
            except Exception:
                pass
        QMessageBox.information(w, "XLSX Export", f"XLSX exported successfully:\n{path}")

    def export_printable_summary(self) -> None:
        w = self.window
        start_date, end_date = w._iso_range_from_edits(w.report_from_date, w.report_to_date)
        path, _ = QFileDialog.getSaveFileName(
            w,
            "Save Printable Summary",
            f"summary_{w._range_suffix(start_date, end_date)}.txt",
            "Text Files (*.txt)",
        )
        if not path:
            return

        lines = [
            "Cafe POS - Printable Summary",
            f"Range: {start_date} to {end_date}",
            "",
            f"Sales: {w.sales_value.text()}",
            f"COGS: {w.cogs_value.text()}",
            f"Gross Profit: {w.gross_profit_value.text()}",
            f"Purchases: {w.purchases_value.text()}",
            f"Expenses: {w.expenses_value.text()}",
            f"Period Fixed Cost: {w.fixed_daily_value.text()}",
            f"Net Profit: {w.net_value.text()}",
            f"Waste Summary: {w.waste_summary_label.text() if hasattr(w, 'waste_summary_label') else 'N/A'}",
            "",
            "Top Selling Items:",
        ]
        for row in range(min(w.top_items_table.rowCount(), 15)):
            item = w.top_items_table.item(row, 0).text() if w.top_items_table.item(row, 0) else ""
            qty = w.top_items_table.item(row, 1).text() if w.top_items_table.item(row, 1) else ""
            value = w.top_items_table.item(row, 2).text() if w.top_items_table.item(row, 2) else ""
            lines.append(f"- {item}: qty {qty}, value {value}")

        try:
            Path(path).write_text("\n".join(lines), encoding="utf-8")
        except Exception as exc:
            QMessageBox.critical(w, "Printable Summary", f"Failed: {exc}")
            return

        self.log_audit("printable_summary_export", "report", "range", f"{start_date}..{end_date}")
        if w.open_after_export_checkbox.isChecked() and hasattr(os, "startfile"):
            try:
                os.startfile(path)
            except Exception:
                pass
        QMessageBox.information(w, "Printable Summary", f"Summary exported:\n{path}")

    def export_audit_csv(self) -> None:
        w = self.window
        if not hasattr(w, "audit_log_table"):
            return
        rows = w._table_rows_for_csv(w.audit_log_table)
        w._save_csv_rows("audit_log_export.csv", rows)
        self.log_audit("audit_export_csv", "audit_log", "table", "audit log exported as csv")

    def export_audit_xlsx(self) -> None:
        w = self.window
        try:
            from openpyxl import Workbook
        except Exception:
            QMessageBox.warning(w, "Audit XLSX", "openpyxl is not installed. Run: pip install openpyxl")
            return

        path, _ = QFileDialog.getSaveFileName(
            w,
            "Save Audit XLSX",
            "audit_log_export.xlsx",
            "Excel Workbook (*.xlsx)",
        )
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "AuditLog"
        for row in w._table_rows_for_csv(w.audit_log_table):
            ws.append(row)

        try:
            wb.save(path)
        except Exception as exc:
            QMessageBox.critical(w, "Audit XLSX", f"Failed to save XLSX: {exc}")
            return

        self.log_audit("audit_export_xlsx", "audit_log", "table", "audit log exported as xlsx")
        if w.open_after_export_checkbox.isChecked() and hasattr(os, "startfile"):
            try:
                os.startfile(path)
            except Exception:
                pass
        QMessageBox.information(w, "Audit XLSX", f"Audit XLSX exported successfully:\n{path}")

    def close_day(self) -> None:
        w = self.window
        pin = self.require_admin_access("Close Day")
        if pin is None:
            return

        selected_date, ok = QInputDialog.getText(
            w,
            "Close Day",
            "Enter date to close (YYYY-MM-DD):",
            text=date.today().isoformat(),
        )
        if not ok:
            return

        try:
            result = w.bookkeeping_service.close_day(selected_date.strip())
        except ValueError as exc:
            QMessageBox.warning(w, "Close Day", str(exc))
            return
        except Exception as exc:
            QMessageBox.critical(w, "Close Day Error", str(exc))
            return

        QMessageBox.information(
            w,
            "Day Closed",
            (
                f"Date: {result['closure_date']}\n"
                f"Sales: INR {result['sales']:.2f}\n"
                f"COGS: INR {result['cogs']:.2f}\n"
                f"Expenses: INR {result['expenses']:.2f}\n"
                f"Gross: INR {result['gross_profit']:.2f}\n"
                f"Net: INR {result['net_profit']:.2f}"
            ),
        )
        self.log_audit("day_close", "closure", result["closure_date"], "manual close-day")

    def backup_now(self) -> None:
        w = self.window
        try:
            backup_path = create_backup(db_path=w.db_path)
        except Exception as exc:
            QMessageBox.critical(w, "Backup Failed", str(exc))
            return

        self.log_audit("backup_now", "backup", str(backup_path), "manual backup")
        QMessageBox.information(w, "Backup Created", f"Backup saved at:\n{backup_path}")

    def export_backup_dialog(self) -> None:
        w = self.window
        folder = QFileDialog.getExistingDirectory(w, "Select Export Folder")
        if not folder:
            return

        try:
            path = export_backup(destination_dir=folder, db_path=w.db_path)
        except Exception as exc:
            QMessageBox.critical(w, "Export Failed", str(exc))
            return

        self.log_audit("backup_export", "backup", str(path), f"export folder={folder}")
        QMessageBox.information(w, "Export Complete", f"Backup exported to:\n{path}")

    def restore_backup_dialog(self) -> None:
        w = self.window
        file_path, _ = QFileDialog.getOpenFileName(w, "Select Backup File", filter="DB Files (*.db)")
        if not file_path:
            return

        pin = self.require_admin_access("Restore Backup")
        if pin is None:
            return

        try:
            current_counts = w.bookkeeping_service.current_database_counts()
            backup_counts = inspect_backup_counts(file_path)
            preview = self.format_restore_preview(
                current_counts=current_counts,
                backup_counts=backup_counts,
                backup_file=file_path,
            )
        except Exception as exc:
            QMessageBox.critical(w, "Restore Pre-check Failed", str(exc))
            return

        confirm = QMessageBox.question(
            w,
            "Confirm Restore",
            preview,
            QMessageBox.Yes | QMessageBox.No,
        )
        if confirm != QMessageBox.Yes:
            return

        try:
            restore_backup(backup_file=file_path, db_path=w.db_path)
        except Exception as exc:
            QMessageBox.critical(w, "Restore Failed", str(exc))
            return

        QMessageBox.information(
            w,
            "Restore Complete",
            "Backup restored. Please restart the application to reload all data safely.",
        )
        self.log_audit("backup_restore", "backup", file_path, "restore completed")

    def save_pending_cart(self) -> None:
        w = self.window
        w.cart_file.parent.mkdir(parents=True, exist_ok=True)
        if not w.cart:
            w.cart_file.unlink(missing_ok=True)
            return

        payload = {"items": list(w.cart.values())}
        w.cart_file.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    def load_pending_cart(self) -> None:
        w = self.window
        if not w.cart_file.exists():
            return

        try:
            payload = json.loads(w.cart_file.read_text(encoding="utf-8"))
            items = payload.get("items", [])
        except Exception:
            w.cart_file.unlink(missing_ok=True)
            return

        if not items:
            w.cart_file.unlink(missing_ok=True)
            return

        answer = QMessageBox.question(
            w,
            "Recover Pending Bill",
            "Found an unsaved cart from previous session. Restore it?",
            QMessageBox.Yes | QMessageBox.No,
        )
        if answer != QMessageBox.Yes:
            w.cart_file.unlink(missing_ok=True)
            return

        w.cart = {int(item["item_id"]): item for item in items}
        w.refresh_cart_table()
